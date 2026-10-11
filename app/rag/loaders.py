"""
文档解析层

把不同格式的文件统一解析成「带结构信息的纯文本」。这里刻意不做向量化和切块,
只解决一个问题:怎么把 PDF / Word / Markdown / 纯文本里的人类排版还原成
对检索友好的文本流。

设计要点:
1. 每个格式一个 loader,统一签名 -> 用注册表模式扩展新格式,新增格式不用改主流程;
2. loader 解析失败不抛出中断整批入库,而是返回 warnings,让上层决定跳过还是继续;
3. 保留 heading 信息(Markdown 标题、Word 样式标题),切块时用于补充「这一段在讲什么」。
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path
from typing import Callable

from app.rag.ocr import OcrEngine, OcrOptions, default_ocr_engine
from app.rag.render import PdfRenderer, build_pdf_renderer

# PDF 页渲染倍率:OCR 对分辨率敏感,2x 是质量与耗时的折中(实测 2x 比 1x 多认出正文小字)
OCR_RENDER_SCALE = 2.0

# 图片后缀:这些文件没有文本层,唯一的取文本路径是 OCR(app/rag/ocr.py)
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".bmp", ".webp", ".tif", ".tiff"}

# 支持的文件后缀 -> loader 名称
SUPPORTED_SUFFIXES = {".txt", ".md", ".markdown", ".pdf", ".docx", ".csv"} | IMAGE_SUFFIXES


@dataclass
class LoadedDocument:
    """一份解析完成的文档。

    :param text: 归一化后的正文(段落之间用空行分隔)
    :param source_path: 相对路径,作为引用与去重依据
    :param title: 文档标题(优先取 Markdown 一级标题 / 文件名)
    :param headings: 标题标记位置,元素为 (字符偏移, 标题文本),切块时据此补结构上下文
    :param warnings: 解析过程中的非致命问题(如 PDF 部分页无文本层)
    :param extraction: 文本来源,取值 ``text_layer`` / ``ocr`` / ``mixed`` / ``none``
        - ``text_layer``:全部来自文档自带的文本层(**未调用 OCR**);
        - ``ocr``:全部来自 OCR(图片文件,或整页扫描件);
        - ``mixed``:部分页用文本层、部分页用 OCR(真实文档最常见:正文有文本层、封面/图表页是图);
        - ``none``:没有取到任何文本(OCR 关闭或后端不可用)。

        引用与排障时据此区分「这句话是原文还是识别出来的」——OCR 有错字(见 docs/rag-ocr.md),
        标注来源后才能判断要不要人工核对。
    """

    text: str
    source_path: str
    title: str = ""
    headings: list[tuple[int, str]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    extraction: str = "text_layer"

    @property
    def char_count(self) -> int:
        return len(self.text)


# ────────────────────────────── 通用文本处理 ──────────────────────────────


def normalize_text(text: str) -> str:
    """归一化文本:统一换行、去掉零宽字符、压缩超长空行。

    PDF 抽取出来的文本常带 \\xa0(不换行空格)、零宽空格和大量连续空行,
    不清理会污染 BM25 分词与 chunk 边界。
    """
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\xa0", " ").replace("\u200b", "").replace("\ufeff", "")
    # 中文字符之间的多余空格(PDF 逐字抽取的常见副作用)。
    # 注意必须用 [ \t]+ 而不是 \s+:后者包含换行,会把中文之间的换行一并删掉,
    # 于是「# 标题\n\n正文」被压成「# 标题正文」——标题吞掉正文 → 标题识别失败 →
    # 整篇退化成 1 个 chunk、heading 全空。这种退化不报错、只降质,极难发现
    # (与 re.MULTILINE 那个坑症状相同但根因不同,详见 docs/rag-self-built.md)。
    text = re.sub(r"(?<=[\u4e00-\u9fff])[ \t]+(?=[\u4e00-\u9fff])", "", text)
    text = re.sub(r"[ \t]{2,}", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def read_text_file(path: Path) -> str:
    """读取文本文件,自动处理常见中文编码。

    中文 Windows 上不少 txt 是 GBK 编码,直接用 utf-8 读会抛 UnicodeDecodeError;
    这里按 utf-8-sig -> utf-8 -> gbk -> gb18030 顺序尝试,最后用 errors="replace" 兜底。
    """
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gbk", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


_MD_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*$", re.MULTILINE)


def extract_markdown_headings(text: str) -> list[tuple[int, str]]:
    """抽取 Markdown 标题及其在文本中的偏移量。"""
    return [(m.start(), m.group(2)) for m in _MD_HEADING_RE.finditer(text)]


# ────────────────────────────── 各格式 loader ──────────────────────────────


def load_plain_text(path: Path, source_path: str) -> LoadedDocument:
    """纯文本:整体作为正文,无结构信息。"""
    text = normalize_text(read_text_file(path))
    return LoadedDocument(text=text, source_path=source_path, title=path.stem)


def load_markdown(path: Path, source_path: str) -> LoadedDocument:
    """Markdown:保留标题层级,标题本身留在正文里作为天然语义分隔。"""
    text = normalize_text(read_text_file(path))
    headings = extract_markdown_headings(text)
    title = headings[0][1] if headings else path.stem
    warnings: list[str] = []
    if not text:
        warnings.append("文件内容为空")
    return LoadedDocument(
        text=text, source_path=source_path, title=title, headings=headings, warnings=warnings
    )


def _open_pdf_reader(path: Path):
    """打开 PDF;pip 不可用的环境里回落到项目根 vendor/ 下的 pypdf(见 app/rag/vendor.py)。"""
    try:
        from pypdf import PdfReader
    except ImportError:
        from app.rag.vendor import ensure_vendor_path

        ensure_vendor_path()
        try:
            from pypdf import PdfReader
        except ImportError as exc:  # pragma: no cover - 依赖确实缺失时给出可执行提示
            raise RuntimeError("缺少 pypdf 依赖,请执行: pip install pypdf") from exc
    return PdfReader(str(path))


@lru_cache(maxsize=4)
def _default_renderer(backend: str = "windows") -> PdfRenderer:
    """进程内共享的 PDF 渲染器(构造时会探测 PowerShell,缓存避免逐页探测)。"""
    return build_pdf_renderer(backend)


def load_pdf(
    path: Path,
    source_path: str,
    *,
    ocr: OcrOptions | None = None,
    ocr_engine: OcrEngine | None = None,
    renderer: PdfRenderer | None = None,
) -> LoadedDocument:
    """PDF:先逐页抽取文本层,**文本层过少的页自动走 OCR 兜底**。

    为什么按页判定而不是整份判定:真实文档常见「正文有文本层、封面/图表页是扫描图」。
    整份判定要么让这种混合文档丢掉图里的内容,要么对整份做无谓且昂贵的 OCR。

    OCR 需要同时满足三个前提,缺一就**安静降级为纯文本层**并留 warning(不抛异常):
    1. `ocr.enabled` 为真且 `min_text_chars > 0`;
    2. OCR 后端可用;
    3. 有可用的 PDF 渲染器(Windows 内置 `Windows.Data.Pdf`)。

    :param ocr: OCR 选项;None 表示从项目 Settings 读取
    :param ocr_engine: 注入的 OCR 后端(测试用);None 表示按配置构造
    :param renderer: 注入的 PDF 渲染器(测试用);None 表示按配置构造
    """
    reader = _open_pdf_reader(path)
    opts = ocr or OcrOptions.from_settings()

    parts: list[str] = []
    headings: list[tuple[int, str]] = []
    warnings: list[str] = []
    empty_pages: list[int] = []
    ocr_pages: list[int] = []
    text_pages: list[int] = []
    offset = 0

    engine: OcrEngine | None = None
    page_renderer: PdfRenderer | None = None
    ocr_ready = False
    if opts.enabled and opts.min_text_chars > 0:
        engine = ocr_engine if ocr_engine is not None else default_ocr_engine(opts)
        page_renderer = renderer if renderer is not None else _default_renderer()
        if engine is None or not engine.available:
            warnings.append(
                f"OCR 后端不可用(backend={opts.backend}),扫描件页将取不到内容;"
                "可用 RAG_OCR_BACKEND 指定 windows/rapidocr,或确认 Windows 已装对应语言包"
            )
        elif page_renderer is None or not page_renderer.available:
            warnings.append("缺少可用的 PDF 渲染器,扫描件页将取不到内容")
        else:
            ocr_ready = True

    # 只有真要 OCR 时才把整份 PDF 读进内存
    pdf_bytes = path.read_bytes() if ocr_ready else b""

    for page_no, page in enumerate(reader.pages, start=1):
        try:
            page_text = normalize_text(page.extract_text() or "")
        except Exception as exc:  # 单页解析失败不应毁掉整份文档
            page_text = ""
            warnings.append(f"第 {page_no} 页解析失败:{exc}")

        used_ocr = False
        if ocr_ready and len(page_text) < opts.min_text_chars:
            ocr_text = _ocr_pdf_page(pdf_bytes, page_no, engine, page_renderer, warnings)
            if ocr_text:
                # 文本层本来就少 → OCR 给得更多就以它为主;两者都有就都留:
                # 宁可多留一点(切块与去重会处理),也不要因为「择一」而丢字。
                if not page_text:
                    page_text = ocr_text
                elif len(ocr_text) > len(page_text):
                    page_text = ocr_text + "\n" + page_text
                else:
                    page_text = page_text + "\n" + ocr_text
                used_ocr = True

        if not page_text:
            empty_pages.append(page_no)
            continue
        # 用页码作为伪标题,让引用能定位到「第 N 页」
        headings.append((offset, f"第 {page_no} 页"))
        parts.append(page_text)
        offset += len(page_text) + 2
        (ocr_pages if used_ocr else text_pages).append(page_no)

    if empty_pages:
        warnings.append(
            f"共 {len(empty_pages)} 页未提取到文本(可能是扫描件/图片页):{empty_pages[:10]}"
        )
    if not parts:
        if ocr_ready:
            warnings.append("整份 PDF 未提取到任何文本,OCR 也未识别出内容")
        else:
            warnings.append("整份 PDF 未提取到任何文本,且 OCR 不可用或未启用")

    if ocr_pages and text_pages:
        extraction = "mixed"
    elif ocr_pages:
        extraction = "ocr"
    elif text_pages:
        extraction = "text_layer"
    else:
        extraction = "none"

    return LoadedDocument(
        text="\n\n".join(parts),
        source_path=source_path,
        title=path.stem,
        headings=headings,
        warnings=warnings,
        extraction=extraction,
    )


def _ocr_pdf_page(
    pdf_bytes: bytes,
    page_no: int,
    engine: OcrEngine,
    renderer: PdfRenderer,
    warnings: list[str],
) -> str:
    """渲染第 page_no 页(1-based)并 OCR;任何失败都只留 warning 并返回空串。"""
    try:
        png = renderer.render_page(pdf_bytes, page_no - 1, scale=OCR_RENDER_SCALE)
    except Exception as exc:
        warnings.append(f"第 {page_no} 页渲染失败,跳过 OCR:{exc}")
        return ""
    if not png:
        warnings.append(f"第 {page_no} 页渲染结果为空,跳过 OCR")
        return ""
    try:
        text = normalize_text(engine.recognize(png))
    except Exception as exc:
        # 第三方 OCR 引擎可能抛异常(接口约定是降级不是中断),这里兜住
        # 单页识别失败不应毁掉整批入库,只留 warning 并跳过该页 OCR。
        warnings.append(f"第 {page_no} 页 OCR 识别失败,跳过:{exc}")
        return ""
    if not text:
        warnings.append(f"第 {page_no} 页 OCR 未识别出文字")
    return text


def load_image(
    path: Path,
    source_path: str,
    *,
    ocr: OcrOptions | None = None,
    ocr_engine: OcrEngine | None = None,
) -> LoadedDocument:
    """图片:整张图走 OCR 变文本(图片没有文本层,这是唯一路径)。

    OCR 关闭或后端不可用时**不抛异常**:返回空文本 + warning,由 pipeline 按
    「未解析出文本」跳过该文件(与既有 loader 约定一致)。
    """
    opts = ocr or OcrOptions.from_settings()
    warnings: list[str] = []
    if not opts.enabled:
        return LoadedDocument(
            text="",
            source_path=source_path,
            title=path.stem,
            warnings=["OCR 已关闭(RAG_OCR_ENABLED=false),图片无法解析"],
            extraction="none",
        )
    engine = ocr_engine if ocr_engine is not None else default_ocr_engine(opts)
    if engine is None or not engine.available:
        return LoadedDocument(
            text="",
            source_path=source_path,
            title=path.stem,
            warnings=[f"OCR 后端不可用(backend={opts.backend}),图片无法解析"],
            extraction="none",
        )
    try:
        raw = path.read_bytes()
    except OSError as exc:
        return LoadedDocument(
            text="",
            source_path=source_path,
            title=path.stem,
            warnings=[f"图片读取失败:{exc}"],
            extraction="none",
        )
    try:
        text = normalize_text(engine.recognize(raw))
    except Exception as exc:
        return LoadedDocument(
            text="",
            source_path=source_path,
            title=path.stem,
            warnings=[f"图片 OCR 识别失败:{exc}"],
            extraction="none",
        )
    if not text:
        warnings.append("OCR 未从图片中识别出文字")
    return LoadedDocument(
        text=text,
        source_path=source_path,
        title=path.stem,
        warnings=warnings,
        extraction="ocr" if text else "none",
    )


def load_docx(path: Path, source_path: str) -> LoadedDocument:
    """Word:按段落读取,识别 Heading 样式作为标题。

    表格同样抽取为 Markdown 表格:企业制度/预算类文档里关键信息常在表里,
    丢掉表格会让 RAG「查不到数字」。
    """
    try:
        import docx  # python-docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("缺少 python-docx 依赖,请执行: pip install python-docx") from exc

    document = docx.Document(str(path))
    blocks: list[str] = []
    headings: list[tuple[int, str]] = []
    warnings: list[str] = []
    offset = 0

    def append_block(content: str) -> None:
        nonlocal offset
        blocks.append(content)
        offset += len(content) + 2

    def iter_blocks(doc):
        """按文档顺序同时遍历段落与表格(python-docx 的 document.paragraphs 会漏掉表格)。"""
        from docx.oxml.ns import qn

        body = doc.element.body
        for child in body.iterchildren():
            if child.tag == qn("w:p"):
                yield Paragraph(child, doc)
            elif child.tag == qn("w:tbl"):
                yield Table(child, doc)

    for block in iter_blocks(document):
        if isinstance(block, Table):
            rows = []
            for row in block.rows:
                cells = [normalize_text(cell.text) for cell in row.cells]
                rows.append("| " + " | ".join(cells) + " |")
            if rows:
                # 分隔行让 Markdown 表格语义完整,也方便后续按行切块
                rows.insert(1, "| " + " | ".join(["---"] * len(block.rows[0].cells)) + " |")
                append_block("\n".join(rows))
            continue

        text = normalize_text(block.text)
        if not text:
            continue
        style_name = (block.style.name or "") if block.style is not None else ""
        if style_name.lower().startswith("heading") or style_name.startswith("标题"):
            # 把标题写成 Markdown 形式,切块器就能复用同一套结构识别逻辑
            level = 2
            digits = re.findall(r"\d+", style_name)
            if digits:
                level = min(max(int(digits[0]), 1), 6)
            headings.append((offset, text))
            append_block("#" * level + " " + text)
        else:
            append_block(text)

    text = normalize_text("\n\n".join(blocks))
    if not text:
        warnings.append("文档正文为空")
    return LoadedDocument(
        text=text, source_path=source_path, title=path.stem, headings=headings, warnings=warnings
    )


def load_csv(path: Path, source_path: str) -> LoadedDocument:
    """CSV:逐行转成「表头: 值」的句子。

    结构化表格逐行成句后,BM25 与向量都能匹配到「某个字段等于某个值」的问法。
    这是 MVP 的简化做法,真正的表格问答应该走数据库/SQL 或表格专用检索。
    """
    text = read_text_file(path)
    reader = csv.reader(io.StringIO(text))
    rows = [row for row in reader if any(cell.strip() for cell in row)]
    if not rows:
        return LoadedDocument(text="", source_path=source_path, title=path.stem, warnings=["CSV 为空"])

    header = [h.strip() or f"列{i + 1}" for i, h in enumerate(rows[0])]
    lines: list[str] = [f"# {path.stem}"]
    for row in rows[1:]:
        pairs = [
            f"{header[i]}: {value.strip()}"
            for i, value in enumerate(row)
            if i < len(header) and value.strip()
        ]
        if pairs:
            lines.append("；".join(pairs) + "。")
    return LoadedDocument(
        text=normalize_text("\n\n".join(lines)), source_path=source_path, title=path.stem
    )


# 后缀 -> loader 注册表
_LOADERS: dict[str, Callable[..., LoadedDocument]] = {
    ".txt": load_plain_text,
    ".md": load_markdown,
    ".markdown": load_markdown,
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".csv": load_csv,
}
# 图片格式全部走同一个 OCR loader
_LOADERS.update({suffix: load_image for suffix in IMAGE_SUFFIXES})

# 需要 OCR 相关参数的 loader(其余 loader 保持原签名,不被迫感知 OCR)
_OCR_AWARE_SUFFIXES = {".pdf"} | IMAGE_SUFFIXES


def load_document(
    path: Path,
    base_dir: Path | None = None,
    *,
    ocr: OcrOptions | None = None,
    ocr_engine: OcrEngine | None = None,
    renderer: PdfRenderer | None = None,
) -> LoadedDocument:
    """按后缀分派 loader 并解析单个文件。

    :param path: 待解析文件
    :param base_dir: 计算 source_path 的基准目录,便于引用时只显示相对路径
    :param ocr: OCR 选项;None 表示从项目 Settings 读取
    :param ocr_engine: 注入的 OCR 后端(测试用;None=按配置构造)
    :param renderer: 注入的 PDF 渲染器(仅 PDF 用;测试用)
    """
    path = Path(path)
    suffix = path.suffix.lower()
    if not path.exists():
        raise FileNotFoundError(f"文件不存在:{path}")
    if suffix not in _LOADERS:
        raise ValueError(f"暂不支持的文件类型 {suffix},支持:{sorted(_LOADERS)}")

    try:
        source_path = str(path.relative_to(base_dir)).replace("\\", "/") if base_dir else path.name
    except ValueError:
        source_path = path.name

    loader = _LOADERS[suffix]
    if suffix == ".pdf":
        return loader(path, source_path, ocr=ocr, ocr_engine=ocr_engine, renderer=renderer)
    if suffix in IMAGE_SUFFIXES:
        return loader(path, source_path, ocr=ocr, ocr_engine=ocr_engine)
    return loader(path, source_path)


def iter_supported_files(root: Path) -> list[Path]:
    """递归收集目录下所有支持的文件(跳过隐藏文件与临时文件)。"""
    root = Path(root)
    if root.is_file():
        return [root] if root.suffix.lower() in _LOADERS else []
    files: list[Path] = []
    for item in sorted(root.rglob("*")):
        if not item.is_file():
            continue
        if item.name.startswith(".") or item.name.startswith("~$"):
            continue
        if item.suffix.lower() in _LOADERS:
            files.append(item)
    return files


__all__ = [
    "LoadedDocument",
    "SUPPORTED_SUFFIXES",
    "IMAGE_SUFFIXES",
    "load_document",
    "iter_supported_files",
    "normalize_text",
    "read_text_file",
    "load_image",
    "load_pdf",
]
